import json
import os
import sys
import datetime
import tkinter as tk
from tkinter import ttk, messagebox, filedialog
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

TYPES = ["Плановая", "Неплановая", "Аварийная", "Технологическое нарушение"]
STATUSES = ["Открыта", "В работе", "Закрыта"]
DT_FMT = "%d.%m.%Y %H:%M"
DT_HINT = "ДД.ММ.ГГГГ ЧЧ:ММ, например 14.10.2026 08:30"

FIELDS = [
    ("date", "Дата"),
    ("type", "Тип работ"),
    ("object", "Объект"),
    ("equipment", "Оборудование"),
    ("description", "Описание"),
    ("works", "Какие работы выполнены"),
    ("cause", "Причина отключения"),
    ("responsible", "Ответственный"),
    ("opened", "Открытие заявки"),
    ("closed", "Закрытие заявки"),
    ("off", "Отключение"),
    ("on", "Включение"),
    ("duration", "Длительность отключения, мин"),
    ("status", "Статус"),
]
LABELS = dict(FIELDS)
INPUT_KEYS = [k for k, _ in FIELDS if k != "duration"]
DT_KEYS = ("opened", "closed", "off", "on")
WIDTHS = {"date": 90, "type": 170, "object": 120, "equipment": 120, "description": 240,
          "works": 200, "cause": 180, "responsible": 120, "opened": 130, "closed": 130,
          "off": 130, "on": 130, "duration": 150, "status": 90}


def data_path():
    if getattr(sys, "frozen", False):
        base = os.path.dirname(sys.executable)
    else:
        base = os.path.dirname(os.path.abspath(__file__))
    return os.path.join(base, "zayavki.json")


def today():
    return datetime.date.today().strftime("%d.%m.%Y")


def parse_dt(s):
    try:
        return datetime.datetime.strptime(s.strip(), DT_FMT)
    except (ValueError, AttributeError):
        return None


def minutes_between(off, on):
    t1, t2 = parse_dt(off), parse_dt(on)
    if t1 is None or t2 is None:
        return ""
    return int((t2 - t1).total_seconds() // 60)


def dur(r):
    v = r.get("duration", "")
    return v if isinstance(v, int) else 0


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Заявки на работы")
        self.geometry("1280x760")
        self.records = self.load()
        self.build_ui()
        self.refresh()

    def load(self):
        try:
            with open(data_path(), encoding="utf-8") as f:
                return json.load(f)
        except (OSError, ValueError):
            return []

    def save(self):
        with open(data_path(), "w", encoding="utf-8") as f:
            json.dump(self.records, f, ensure_ascii=False, indent=2)

    def build_ui(self):
        form = ttk.LabelFrame(self, text="Новая заявка", padding=10)
        form.pack(fill="x", padx=10, pady=8)
        self.vars = {}
        for i, key in enumerate(INPUT_KEYS):
            r, c = divmod(i, 3)
            ttk.Label(form, text=LABELS[key]).grid(row=r * 2, column=c, sticky="w", padx=6)
            if key == "type":
                w = ttk.Combobox(form, values=TYPES, state="readonly")
            elif key == "status":
                w = ttk.Combobox(form, values=STATUSES, state="readonly")
            else:
                w = ttk.Entry(form)
            w.grid(row=r * 2 + 1, column=c, sticky="ew", padx=6, pady=(0, 6))
            self.vars[key] = w
        for c in range(3):
            form.columnconfigure(c, weight=1)
        self.vars["type"].set(TYPES[0])
        self.vars["status"].set(STATUSES[0])
        self.clear()

        rows = (len(INPUT_KEYS) + 2) // 3
        ttk.Label(form, text=f"Формат времени: {DT_HINT}").grid(
            row=rows * 2, column=0, columnspan=3, sticky="w", padx=6)
        btns = ttk.Frame(form)
        btns.grid(row=rows * 2 + 1, column=0, columnspan=3, sticky="e", pady=6)
        ttk.Button(btns, text="Добавить", command=self.add).pack(side="left", padx=4)
        ttk.Button(btns, text="Очистить поля", command=self.clear).pack(side="left", padx=4)

        bar = ttk.Frame(self)
        bar.pack(fill="x", padx=10)
        ttk.Button(bar, text="Сменить статус выбранной", command=self.change_status).pack(side="left")
        ttk.Button(bar, text="Удалить выбранную", command=self.delete).pack(side="left", padx=6)
        ttk.Button(bar, text="Выгрузить в Excel", command=self.export).pack(side="right")

        cols = [k for k, _ in FIELDS]
        frame = ttk.Frame(self)
        frame.pack(fill="both", expand=True, padx=10, pady=8)
        self.tree = ttk.Treeview(frame, columns=cols, show="headings", selectmode="browse")
        for k, _ in FIELDS:
            self.tree.heading(k, text=LABELS[k])
            self.tree.column(k, width=WIDTHS[k], anchor="w")
        ysb = ttk.Scrollbar(frame, orient="vertical", command=self.tree.yview)
        xsb = ttk.Scrollbar(frame, orient="horizontal", command=self.tree.xview)
        self.tree.configure(yscrollcommand=ysb.set, xscrollcommand=xsb.set)
        self.tree.grid(row=0, column=0, sticky="nsew")
        ysb.grid(row=0, column=1, sticky="ns")
        xsb.grid(row=1, column=0, sticky="ew")
        frame.rowconfigure(0, weight=1)
        frame.columnconfigure(0, weight=1)

    def clear(self):
        for w in self.vars.values():
            if isinstance(w, ttk.Entry):
                w.delete(0, "end")
        self.vars["date"].insert(0, today())

    def add(self):
        rec = {k: self.vars[k].get().strip() for k in INPUT_KEYS}
        if not rec["object"] or not rec["description"]:
            messagebox.showwarning("Нет данных", "Заполните поля «Объект» и «Описание».")
            return
        for k in DT_KEYS:
            if rec[k] and parse_dt(rec[k]) is None:
                messagebox.showwarning("Дата и время",
                                       f"Поле «{LABELS[k]}» заполните в формате:\n{DT_HINT}")
                return
        if rec["off"] and rec["on"] and parse_dt(rec["on"]) <= parse_dt(rec["off"]):
            messagebox.showwarning("Время", "Время включения должно быть позже времени отключения.")
            return
        rec["duration"] = minutes_between(rec["off"], rec["on"])
        self.records.append(rec)
        self.save()
        self.refresh()
        self.clear()

    def selected_index(self):
        sel = self.tree.selection()
        if not sel:
            messagebox.showinfo("Выбор", "Сначала выберите заявку в таблице.")
            return None
        return self.tree.index(sel[0])

    def change_status(self):
        i = self.selected_index()
        if i is None:
            return
        cur = self.records[i].get("status", "")
        nxt = STATUSES[(STATUSES.index(cur) + 1) % len(STATUSES)] if cur in STATUSES else STATUSES[0]
        self.records[i]["status"] = nxt
        self.save()
        self.refresh()

    def delete(self):
        i = self.selected_index()
        if i is None:
            return
        if messagebox.askyesno("Удаление", "Удалить выбранную заявку?"):
            del self.records[i]
            self.save()
            self.refresh()

    def refresh(self):
        self.tree.delete(*self.tree.get_children())
        for r in self.records:
            self.tree.insert("", "end", values=[r.get(k, "") for k, _ in FIELDS])

    def export(self):
        if not self.records:
            messagebox.showinfo("Пусто", "Нет заявок для выгрузки.")
            return
        path = filedialog.asksaveasfilename(
            defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")],
            initialfile=f"zayavki_{datetime.date.today():%Y-%m-%d}.xlsx")
        if not path:
            return

        wb = Workbook()
        ws = wb.active
        ws.title = "Заявки"
        headers = ["№"] + [label for _, label in FIELDS]
        ws.append(headers)
        for n, r in enumerate(self.records, 1):
            row = [n]
            for k, _ in FIELDS:
                v = r.get(k, "")
                if k in DT_KEYS:
                    v = parse_dt(v) or v
                row.append(v)
            ws.append(row)

        dt_cols = [headers.index(LABELS[k]) + 1 for k in DT_KEYS]
        for col in dt_cols:
            for row in range(2, ws.max_row + 1):
                ws.cell(row=row, column=col).number_format = "dd.mm.yyyy hh:mm"

        thin = Side(style="thin", color="999999")
        border = Border(top=thin, bottom=thin, left=thin, right=thin)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
            for cell in row:
                cell.border = border
                if cell.row > 1:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 45
        for col in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(col)].width = 16
        ws.column_dimensions["A"].width = 6
        for k in ("description", "works", "cause"):
            ws.column_dimensions[get_column_letter(headers.index(LABELS[k]) + 1)].width = 40
        for col in dt_cols:
            ws.column_dimensions[get_column_letter(col)].width = 18
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = ws.dimensions

        summary = wb.create_sheet("Сводка")
        summary.append(["Тип работ", "Количество", "Суммарное отключение, мин"])
        for cell in summary[1]:
            cell.font = Font(bold=True)
        for t in TYPES:
            items = [r for r in self.records if r.get("type") == t]
            summary.append([t, len(items), sum(dur(r) for r in items)])
        summary.append(["Всего", len(self.records), sum(dur(r) for r in self.records)])
        summary.column_dimensions["A"].width = 30
        summary.column_dimensions["B"].width = 14
        summary.column_dimensions["C"].width = 26

        try:
            wb.save(path)
        except OSError as e:
            messagebox.showerror("Ошибка", f"Не удалось сохранить файл:\n{e}")
            return
        messagebox.showinfo("Готово", f"Выгружено заявок: {len(self.records)}\n{path}")


if __name__ == "__main__":
    App().mainloop()
